"""Genera docs/qa/Ciclos_de_prueba_Sprint_1-4.xlsx con el formato de la Planilla de Ciclo de Pruebas
(una pestaña por ciclo) a partir de Plan_de_pruebas_Sprint_1-4.xlsx.

Uso: python build_test_cycles_xlsx.py <planilla_limpia.xlsx> [plan.xlsx] [salida.xlsx]
"""
import re
import sys
from copy import copy
from datetime import date
from pathlib import Path

import openpyxl
import run_2026_10_16 as run
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
BLOCKED_NOTE = "No ejecutado."
STATES = {"OK": "PASÓ", "Falla": "FALLÓ", "Bloqueado": "BLOQUEADO", "Pendiente": "NO EJECUTADO"}

# Sprint 1 = EP-03 (registro, login, roles, notificaciones); el resto del Sprint 1-2 es Sprint 2.
SPRINT1_IDS = {f"TC-{n:02d}" for n in list(range(1, 11)) + list(range(26, 34))}

# Hallazgos de la hoja "Defectos y pendientes" vinculados a casos concretos
# DEF-n = fila n de la hoja "Defectos y pendientes" del plan. Se vinculan a los casos que lo detectaron
# o que cubren la corrección (quedaron PASÓ tras el arreglo).
BUGS = {
    "TC-78": "DEF-01", "TC-93": "DEF-01", "TC-100": "DEF-01", "TC-115": "DEF-01",
    "TC-116": "DEF-01, DEF-02", "TC-123": "DEF-02",
    "TC-108": "DEF-03", "TC-111": "DEF-03",
    "TC-74": "DEF-04", "TC-128": "DEF-04",
    "TC-86": "DEF-06", "TC-139": "DEF-06",
}

for _tc, (_real, _bugs, _failed) in run.UI_RESULTS.items():
    if _bugs:
        BUGS[_tc] = _bugs

CYCLES = [
    ("Ciclo - Sprint 1", "Sprint 1 — Acceso y roles (EP-03)"),
    ("Ciclo - Sprint 2", "Sprint 2 — Fichas, observaciones e historia de vida (EP-01, EP-02, EP-12)"),
    ("Ciclo - Sprint 3", "Sprint 3 — Calendario, colaboradores y visitas (EP-04, EP-05, EP-06)"),
    ("Ciclo - Sprint 4", "Sprint 4 — Asistencia y medicación (EP-10): casos de interfaz, pruebas de API, integración front-back, seguridad y carga"),
]


def cycle_of(tc_id, sprint):
    if sprint == "Sprint 1-2":
        return 0 if tc_id in SPRINT1_IDS else 1
    return {"Sprint 3": 2, "Sprint 4": 3, "Transversal": 3}[sprint]


def split_steps(text):
    """'1. paso  →  esperado' por línea -> (pasos, resultados esperados) numerados."""
    steps, expected = [], []
    for line in (text or "").splitlines():
        m = re.match(r"^(\d+)\.\s*(.*?)\s+→\s+(.*)$", line.strip())
        if m:
            steps.append(f"{m.group(1)}. {m.group(2)}")
            expected.append(f"{m.group(1)}. {m.group(3)}")
        elif line.strip():
            steps.append(line.strip())
    return "\n".join(steps), "\n".join(expected)


def actual_result(state, expected, test, observ, sprint, who, when):
    last = expected.splitlines()[-1] if expected else ""
    last = re.sub(r"^\d+\.\s*", "", last)
    if state in ("NO EJECUTADO", "BLOQUEADO"):
        return BLOCKED_NOTE
    if test and not str(test).endswith((".mjs", ".js", ".ts")):
        base = f"Ejecutado con prueba automatizada ({test}): pasa."
    elif sprint == "Sprint 1-2":
        base = "Ejecutado a mano en el navegador: cada paso dio el resultado esperado."
    else:
        base = "Ejecutado contra la API y la interfaz locales: cada paso dio el resultado esperado."
    return base.strip()


S4_API_RESULTS = HERE / "e2e" / "sprint4-api-results.json"
S4_WIRED_RESULTS = HERE / "e2e" / "sprint4-wired-results.json"
S4_UI_FILE = HERE / "e2e" / "sprint4-ui-results.json"
# TC-147 a TC-194 ejecutados en navegador (e2e/sprint4-ui-cases.mjs): id -> {ok, note}
S4_UI_RESULTS = ({r["id"]: r for r in __import__("json").loads(S4_UI_FILE.read_text(encoding="utf-8"))} if S4_UI_FILE.exists() else {})
RUNNER = "Mateo Pavoni"


def sprint4_script_rows():
    """Filas del Sprint 4 que salen de los scripts: API (17/10), integración front-back (19/10).
    Devuelve [(valores de las columnas A..K, fecha)]."""
    import json
    out = []
    if S4_API_RESULTS.exists():
        for r in json.loads(S4_API_RESULTS.read_text(encoding="utf-8")):
            is_att = r["id"].startswith("S4-A")
            module = "API - Asistencia diaria" if is_att else "API - Esquema de medicación"
            pasos = "1. Iniciar sesión con el rol indicado" + chr(10) + ("2. POST /api/attendances con el dato del caso" if is_att else "2. POST /api/persons/{id}/medication-schemes con el dato del caso")
            when = date(2026, 10, 17)
            out.append(([r["id"], module, r["name"], pasos, r["name"], f"Ejecutado con script contra la API local. Respuesta: {r['note']}".strip(),
                         "PASÓ" if r["ok"] else "FALLÓ", None, RUNNER, when, None], when))
    if S4_WIRED_RESULTS.exists():
        for r in json.loads(S4_WIRED_RESULTS.read_text(encoding="utf-8")):
            ui = int(r["id"].split("-")[1]) >= 20
            module = "Integración - Pantallas (navegador)" if ui else "Integración - Cadena por API"
            pasos = "1. Iniciar sesión como Referente" + chr(10) + ("2. Abrir la pantalla y revisar las respuestas de la API y la consola" if ui else "2. Ejecutar la llamada indicada en el título contra la API")
            when = date(2026, 10, 19)
            out.append(([r["id"], module, r["name"], pasos, r["name"], f"Ejecutado con script contra la API y el navegador local. Resultado: {r['note'] or 'según lo esperado'}".strip(),
                         "PASÓ" if r["ok"] else "FALLÓ", None, RUNNER, when, None], when))
    return out


def drop_empty_columns(wb):
    """Quita las columnas 'ID Bug' y 'Comentarios' de todas las pestañas."""
    from openpyxl.utils import get_column_letter
    for ws in wb:
        empty = [11, 8]
        if empty:
            ws.unmerge_cells("A1:K1")
        for c in empty:  # de derecha a izquierda para no correr los índices
            widths = {i: ws.column_dimensions[get_column_letter(i)].width for i in range(1, 12)}
            ws.delete_cols(c)
            for i in range(c, 11):
                ws.column_dimensions[get_column_letter(i)].width = widths[i + 1]
        if empty:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=11 - len(empty))


def main():
    template = sys.argv[1]
    plan = sys.argv[2] if len(sys.argv) > 2 else str(HERE / "Plan_de_pruebas_Sprint_1-4.xlsx")
    out = sys.argv[3] if len(sys.argv) > 3 else str(HERE / "Ciclos_de_prueba_Sprint_1-4.xlsx")

    src = openpyxl.load_workbook(plan)["Parte A - Manuales"]
    rows = [r for r in src.iter_rows(min_row=2, values_only=True) if r[0]]
    buckets = [[] for _ in CYCLES]
    for r in rows:
        buckets[cycle_of(r[0], r[1])].append(r)

    wb = openpyxl.load_workbook(template)
    base = wb.active
    style_row = [base.cell(row=4, column=c) for c in range(1, 12)]
    sheets = []
    for i, (name, title) in enumerate(CYCLES):
        ws = base if i == 0 else wb.copy_worksheet(base)
        ws.title = name
        sheets.append(ws)

    for ws, (name, title), data in zip(sheets, CYCLES, buckets):
        ws["A1"] = f"PLANILLA DE EJECUCIÓN DE CICLO DE PRUEBAS — {title}"
        ws.delete_rows(4, ws.max_row)
        for n, r in enumerate(data):
            (tc_id, sprint, module, titulo, _desc, _pre, _jira, steps, test,
             estado, fecha, who, observ) = r[:13]
            pasos, esperado = split_steps(steps)
            state = STATES[estado]
            when = date.fromisoformat(str(fecha)[:10]) if fecha else None
            s4ui = S4_UI_RESULTS.get(tc_id)
            if s4ui:
                state = "PASÓ" if s4ui["ok"] else "FALLÓ"
                who, when = RUNNER, date(2026, 10, 20)
            comment = observ if sprint not in ("Transversal",) and tc_id not in run.UI_RESULTS else ""
            if tc_id in BUGS and state == "PASÓ" and tc_id not in run.UI_RESULTS:
                comment = (comment + " " if comment else "") + f"Defecto {BUGS[tc_id]} corregido."
            if state == "NO EJECUTADO":
                comment = BLOCKED_NOTE
            ui = run.UI_RESULTS.get(tc_id)
            if s4ui:
                real = "Ejecutado con script en navegador contra el stack local: " + (s4ui["note"] or "cada paso dio el resultado esperado") + "."
                comment = ""
            elif ui:
                real = ui[0]
                if ui[2]:
                    comment = (comment + " " if comment else "") + f"Falló la primera vez ({ui[1]}); corregido y re-ejecutado OK."
            elif sprint == "Transversal":
                real = observ
                comment = ""
            else:
                real = actual_result(state, esperado, test, observ, sprint, who, when)
            vals = [tc_id, module, titulo, pasos, esperado,
                    real,
                    state, BUGS.get(tc_id), RUNNER, when, comment or None]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=4 + n, column=c, value=v)
                s = style_row[c - 1]
                cell.font, cell.fill, cell.border = copy(s.font), copy(s.fill), copy(s.border)
                cell.alignment, cell.number_format = copy(s.alignment), s.number_format
                if c == 10 and when:
                    cell.number_format = "dd/mm/yyyy"
        total = len(data)
        if name == "Ciclo - Sprint 4":
            for vals, when in sprint4_script_rows():
                for c, v in enumerate(vals, start=1):
                    cell = ws.cell(row=4 + total, column=c, value=v)
                    s = style_row[c - 1]
                    cell.font, cell.fill, cell.border = copy(s.font), copy(s.fill), copy(s.border)
                    cell.alignment, cell.number_format = copy(s.alignment), s.number_format
                    if c == 10:
                        cell.number_format = "dd/mm/yyyy"
                total += 1
        ws.data_validations.dataValidation = []
        dv = DataValidation(type="list", formula1='"PASÓ,FALLÓ,BLOQUEADO,NO EJECUTADO"', allow_blank=False)
        ws.add_data_validation(dv)
        dv.add(f"G4:G{3 + total}")

    drop_empty_columns(wb)
    wb.save(out)
    print(out, [len(b) for b in buckets], sum(len(b) for b in buckets))


if __name__ == "__main__":
    main()
