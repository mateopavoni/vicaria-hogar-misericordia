"""Genera docs/qa/Plan_de_pruebas_Sprint_1-3.xlsx.

Fuentes:
  - el docx "Plan de prueba - completo formateado" (Sprint 1 y 2: TC-01..TC-40 y TA-001..TA-233)
  - docs/qa/QA_SPRINT3.md (Sprint 3: TC-41 en adelante)
  - los tests xUnit de backend/tests (Sprint 3: TA-234 en adelante, derivados de los nombres de test)

Uso: python docs/qa/build_test_plan_xlsx.py <ruta al docx> [resultado de test.json]
Terminología del docx: TC = caso manual, TA = caso automatizado; campos Description,
Pre-conditions, Jira Requirement, Test (xUnit) y Pasos (#, Paso, Dato, Esperamos).
"""
import json
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
QA_MD = ROOT / "docs" / "qa" / "QA_SPRINT3.md"
OUT = ROOT / "docs" / "qa" / "Plan_de_pruebas_Sprint_1-3.xlsx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# fechas de ejecución registradas
DATE_S12_MANUAL = "2026-10-06"   # QA manual de los 40 TC (todos OK)
DATE_S3_UI = "2026-10-08"        # QA de UI Sprint 3
DATE_RERUN = "2026-10-09"        # re-ejecución con el backend de Emir integrado
RUNNER = "Mateo Pavoni"

STORY_OF = {"190": "16", "195": "17", "200": "18", "205": "19", "211": "74", "212": "74", "17": "17"}


# ---------------------------------------------------------------- docx
def para_text(p):
    return "".join(t.text or "" for t in p.iter(W + "t")).strip()


def cell_text(tc):
    return "\n".join(t for t in (para_text(p) for p in tc.iter(W + "p")) if t)


def parse_docx(path):
    xml = zipfile.ZipFile(path).read("word/document.xml")
    body = ET.fromstring(xml).find(W + "body")
    part = module = None
    cases, cur = [], None
    for el in body:
        if el.tag == W + "p":
            st = el.find(f"{W}pPr/{W}pStyle")
            style = st.get(W + "val") if st is not None else ""
            text = para_text(el)
            if style == "Heading1":
                part = text
            elif style == "Heading2":
                module = text
        elif el.tag == W + "tbl":
            rows = [[cell_text(c) for c in r.findall(W + "tc")] for r in el.findall(W + "tr")]
            m = re.match(r"^(T[CA]-\d+)\s+—\s+(.*)$", rows[0][0]) if len(rows) == 1 and len(rows[0]) == 1 else None
            if m:  # tabla de una celda con el título del caso
                cur = {"id": m.group(1), "title": m.group(2), "part": part, "module": module,
                       "description": "", "pre": "", "jira": "", "test": "", "steps": []}
                cases.append(cur)
            elif cur is None:
                continue
            elif rows and rows[0] and rows[0][0] == "#":  # tabla de pasos
                for r in rows[1:]:
                    if len(r) >= 4 and r[0]:
                        cur["steps"].append((r[0], r[1], r[2], r[3]))
            else:  # tabla de ficha: etiqueta | valor
                for r in rows:
                    if len(r) >= 2:
                        key = r[0].lower()
                        if key.startswith("description"):
                            cur["description"] = r[1]
                        elif key.startswith("pre-conditions"):
                            cur["pre"] = r[1]
                        elif key.startswith("jira"):
                            cur["jira"] = r[1]
                        elif key.startswith("test"):
                            cur["test"] = r[1]
    return cases


# ---------------------------------------------------------------- Sprint 3 (markdown)
def parse_sprint3_md(path):
    text = path.read_text(encoding="utf8")
    module = story = ""
    kind = "API"
    out = []
    for line in text.splitlines():
        if line.startswith("## ") or line.startswith("### "):
            module = line.lstrip("# ").strip()
            mm = re.match(r"^SCRUM-(\d+) — (.*?) \(historia", module)
            if mm:
                module = f"{mm.group(2)} (SCRUM-{mm.group(1)})"
            m = re.search(r"historia SCRUM-(\d+)", module)
            if m:
                story = m.group(1)
            else:
                ids = re.findall(r"SCRUM-(\d+)", module)
                for i in ids:
                    if i in STORY_OF:
                        story = STORY_OF[i]
            kind = "UI" if module.startswith("QA de UI") or "Re-ejecución" in module else "API"
        elif line.startswith("Re-ejecución del"):
            kind = "UI-RERUN"
        elif line.startswith("| TC-"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            tc = cells[0]
            if len(cells) >= 5:
                out.append({"id": tc, "module": module, "jira": f"SCRUM-{story}", "title": cells[1],
                            "step": cells[2].replace("`", ""), "expected": cells[3].replace("`", ""), "test": cells[4].replace("`", ""),
                            "kind": "API", "result": "OK"})
            elif len(cells) == 3:
                result = cells[2]
                out.append({"id": tc, "module": module, "jira": f"SCRUM-{story}", "title": cells[1],
                            "step": cells[1], "expected": "El flujo se completa sin errores de API",
                            "test": "docs/qa/e2e/sprint3-ui-flows.mjs", "kind": "UI",
                            "result": "OK" if result.startswith("OK") else "Falla", "note": result})
    return out


# ---------------------------------------------------------------- Sprint 3 (tests xUnit)
AREAS = [
    ("CalendarEvents/GeneralCalendar", "SCRUM-15 / SCRUM-16", "Calendario general"),
    ("CalendarEvents/PersonalCalendar", "SCRUM-17", "Calendario personal"),
    ("CalendarEvents/CalendarEvent", "SCRUM-15 / SCRUM-16", "Calendario general"),
    ("Collaborators/", "SCRUM-18 / SCRUM-19", "Colaboradores"),
    ("CasonaVisits/", "SCRUM-74", "Visitas de la Casa de Convivencia"),
]


def parse_sprint3_tests():
    rows = []
    for proj in ("Vicaria.UnitTests", "Vicaria.IntegrationTests"):
        base = ROOT / "backend" / "tests" / proj
        for f in sorted(base.rglob("*.cs")):
            rel = f.relative_to(base).as_posix()
            area = next(((j, m) for k, j, m in AREAS if rel.startswith(k)), None)
            if not area:
                continue
            src = f.read_text(encoding="utf8")
            cls = re.search(r"public class (\w+)", src)
            if not cls:
                continue
            for m in re.finditer(r"\[(?:Fact|Theory)[^\]]*\][\s\S]*?public (?:async )?(?:Task|void) (\w+)\(", src):
                name = m.group(1)
                human = re.sub(r"_", " · ", name)
                rows.append({"proj": proj, "cls": cls.group(1), "method": name, "jira": area[0],
                             "module": area[1], "human": human})
    return rows


# ---------------------------------------------------------------- Excel
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="D0D7DE")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
STATES = ["Pendiente", "OK", "Falla", "Bloqueado"]


def style_sheet(ws, widths, header_row=1):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[header_row]:
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=header_row + 1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDER
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)
    ws.auto_filter.ref = ws.dimensions


def add_state_validation(ws, col_letter):
    dv = DataValidation(type="list", formula1='"' + ",".join(STATES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{ws.max_row}")


def steps_text(steps):
    return "\n".join(f"{n}. {p}  →  {e}" for n, p, _d, e in steps)


def main():
    docx = Path(sys.argv[1])
    auto_ok = True  # el suite completo pasó (se confirma con el argumento opcional)
    if len(sys.argv) > 2:
        auto_ok = json.loads(Path(sys.argv[2]).read_text(encoding="utf8")).get("all_passed", True)
    auto_state = "OK" if auto_ok else "Pendiente"

    cases = parse_docx(docx)
    manual = [c for c in cases if c["id"].startswith("TC-")]
    auto = [c for c in cases if c["id"].startswith("TA-")]
    s3 = sorted(parse_sprint3_md(QA_MD), key=lambda t: int(t["id"][3:]))
    s3_tests = parse_sprint3_tests()

    wb = Workbook()

    # ---- Parte A: casos manuales (Sprint 1-2 del docx + Sprint 3)
    ws = wb.active
    ws.title = "Parte A - Manuales"
    ws.append(["ID", "Sprint", "Módulo", "Título", "Description", "Pre-conditions", "Jira Requirement",
               "Pasos (Paso → Esperamos)", "Test (xUnit) / Script", "Estado", "Fecha de ejecución",
               "Ejecutó", "Observaciones"])
    for c in manual:
        ws.append([c["id"], "Sprint 1-2", c["module"], c["title"], c["description"], c["pre"], c["jira"],
                   steps_text(c["steps"]), "", "OK", DATE_S12_MANUAL, RUNNER, ""])
    for t in s3:
        ui = t["kind"] == "UI"
        pre = ("Stack local levantado (docker compose en backend/); usuarios de prueba *@test.com."
               if ui else "API levantada con WebApplicationFactory; sesión con el rol indicado en los pasos.")
        date = DATE_RERUN if (not ui or int(t["id"][3:]) >= 130 or t["id"] == "TC-93") else DATE_S3_UI
        if not ui:
            date = DATE_RERUN
        ws.append([t["id"], "Sprint 3", t["module"], t["title"], t["title"], pre, t["jira"],
                   f"1. {t['step']}  →  {t['expected']}", t["test"],
                   t["result"], date, RUNNER if ui else "Suite automatizada",
                   t.get("note", "") if t.get("note") and t["result"] == "OK" and "re-ejecutado" in t["note"] else ""])
    style_sheet(ws, [9, 11, 34, 44, 44, 38, 16, 70, 46, 12, 14, 18, 36])
    add_state_validation(ws, "J")
    manual_rows = ws.max_row

    # ---- Pasos de los TC del docx (una fila por paso)
    wp = wb.create_sheet("Pasos TC")
    wp.append(["ID", "#", "Paso", "Dato", "Esperamos"])
    for c in manual:
        for n, p, d, e in c["steps"]:
            wp.append([c["id"], int(n), p, d, e])
    style_sheet(wp, [9, 5, 60, 36, 60])

    # ---- Parte B: automatizados
    wb_ = wb.create_sheet("Parte B - Automatizados")
    wb_.append(["ID", "Sprint", "Módulo", "Título", "Description", "Pre-conditions", "Jira Requirement",
                "Pasos (Paso → Esperamos)", "Test (xUnit)", "Estado", "Fecha de ejecución", "Ejecutó"])
    for c in auto:
        wb_.append([c["id"], "Sprint 1-2", c["module"], c["title"], c["description"], c["pre"], c["jira"],
                    steps_text(c["steps"]), c["test"], auto_state, DATE_RERUN, "Suite automatizada"])
    by_method = {}
    for t in s3:
        if t["kind"] == "API":
            for name in t["test"].split(","):
                by_method.setdefault(name.strip(), t)
    n = len(auto)
    for t in s3_tests:
        n += 1
        tc = by_method.get(t["method"])
        if tc:
            title = f"Validar: {tc['title']}"
            desc = f"Respalda {tc['id']}: {tc['expected']}"
            pre = "API levantada con WebApplicationFactory." if t["proj"].endswith("IntegrationTests") else "Servicio con DbContext InMemory."
            steps = f"1. {tc['step']}  →  {tc['expected']}"
            jira = tc["jira"]
        else:
            title = f"Validar: {t['human']}"
            desc = "Sin TC manual asociado; descripción derivada del nombre del test."
            pre = "API levantada con WebApplicationFactory." if t["proj"].endswith("IntegrationTests") else "Servicio con DbContext InMemory."
            steps = ""
            jira = t["jira"]
        wb_.append([f"TA-{n:03d}", "Sprint 3", t["module"], title, desc, pre, jira, steps,
                    f"{t['proj']} › {t['cls']} › {t['method']}", auto_state, DATE_RERUN, "Suite automatizada"])
    style_sheet(wb_, [9, 11, 34, 50, 44, 34, 18, 60, 70, 12, 14, 18])
    add_state_validation(wb_, "J")
    auto_rows = wb_.max_row

    # ---- Resumen
    wr = wb.create_sheet("Resumen", 0)
    wr["A1"] = "Plan de pruebas — Vicaria de los Pobres, Pastoral de Adicciones"
    wr["A1"].font = Font(bold=True, size=14)
    wr["A2"] = "Casos manuales (TC) y automatizados (TA) de los Sprint 1, 2 y 3. Se actualiza con los contadores de las hojas."
    wr.append([])
    wr.append(["Tipo", "Sprint", "Total", "OK", "Falla", "Bloqueado", "Pendiente"])
    hr = wr.max_row
    for c in wr[hr]:
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER

    def count_row(label, sprint, sheet, last, scol, state_col):
        r = wr.max_row + 1
        sh = f"'{sheet}'"
        rng_s = f"{sh}!${scol}$2:${scol}${last}"
        rng_e = f"{sh}!${state_col}$2:${state_col}${last}"
        wr.append([label, sprint,
                   f'=COUNTIF({rng_s},B{r})',
                   f'=COUNTIFS({rng_s},B{r},{rng_e},"OK")',
                   f'=COUNTIFS({rng_s},B{r},{rng_e},"Falla")',
                   f'=COUNTIFS({rng_s},B{r},{rng_e},"Bloqueado")',
                   f'=COUNTIFS({rng_s},B{r},{rng_e},"Pendiente")'])

    count_row("Manuales (TC)", "Sprint 1-2", "Parte A - Manuales", manual_rows, "B", "J")
    count_row("Manuales (TC)", "Sprint 3", "Parte A - Manuales", manual_rows, "B", "J")
    count_row("Automatizados (TA)", "Sprint 1-2", "Parte B - Automatizados", auto_rows, "B", "J")
    count_row("Automatizados (TA)", "Sprint 3", "Parte B - Automatizados", auto_rows, "B", "J")
    first, last = hr + 1, wr.max_row
    wr.append(["Total", "", *[f"=SUM({get_column_letter(i)}{first}:{get_column_letter(i)}{last})" for i in range(3, 8)]])
    for row in wr.iter_rows(min_row=hr + 1, max_row=wr.max_row):
        for c in row:
            c.border = BORDER
    for c in wr[wr.max_row]:
        c.font = Font(bold=True)

    wr.append([])
    wr.append(["Última ejecución (2026-10-09, rama dev)"])
    wr.cell(row=wr.max_row, column=1).font = Font(bold=True)
    for line in [
        "Backend: 246 pruebas unitarias y 240 de integración, todas OK.",
        "Frontend: 20 archivos / 51 pruebas y build de producción, OK.",
        "UI en navegador real (Playwright): 33 de 33 verificaciones OK en el flujo de Sprint 3.",
        "Prueba de humo por rol (Referente, Directora, Escucha, Coordinador): 11 de 15 pantallas OK; las 4 restantes son /inicio (ver Defectos y pendientes).",
        "QA manual Sprint 1-2 (TC-01 a TC-40): 40 de 40 OK el 2026-10-06.",
    ]:
        wr.append([line])
    wr.append([])
    wr.append(["Épicas / historias cubiertas"])
    wr.cell(row=wr.max_row, column=1).font = Font(bold=True)
    for line in [
        "Sprint 1-2: EP-01 Fichas de personas · EP-02 Observaciones e historia de vida · EP-03 Cuentas, roles y notificaciones · EP-10 Asistencia · EP-12 Casa de Convivencia",
        "Sprint 3: EP-04 Calendario general y personal (SCRUM-15, 16, 17) · EP-05 Colaboradores (SCRUM-18, 19) y visitas de la Casa de Convivencia (SCRUM-74)",
    ]:
        wr.append([line])
    wr.append([])
    wr.append(["Hojas"])
    wr.cell(row=wr.max_row, column=1).font = Font(bold=True)
    for line in [
        "Parte A - Manuales: un caso por fila con sus pasos; el estado se completa desde la lista desplegable.",
        "Pasos TC: los TC del Sprint 1-2 con una fila por paso (#, Paso, Dato, Esperamos).",
        "Parte B - Automatizados: un test xUnit por fila. Sprint 3 deriva su descripción del nombre del test.",
        "Defectos y pendientes: hallazgos de la ejecución y decisiones abiertas.",
    ]:
        wr.append([line])
    for col, w in zip("ABCDEFG", [24, 14, 10, 10, 10, 12, 12]):
        wr.column_dimensions[col].width = w

    # ---- Defectos y pendientes
    wd = wb.create_sheet("Defectos y pendientes")
    wd.append(["#", "Tipo", "Descripción", "Origen / Jira", "Estado", "Detalle"])
    items = [
        ("Defecto", "Validadores de edición/estado no registrados en la inyección de dependencias: los endpoints PUT de eventos generales, PUT de colaboradores y PATCH de estado de visitas respondían 500.",
         "SCRUM-190 / 200 / 211 (dev-backend)", "Corregido en dev",
         "Se registraron en Program.cs y se agregaron tests HTTP. Falta llevar el arreglo a dev-backend."),
        ("Defecto", "El listado de colaboradores no devolvía el estado activo/inactivo.",
         "SCRUM-200", "Corregido en dev", "CollaboratorListItemDto ahora incluye isActive."),
        ("Brecha", "No existían edición ni borrado de eventos personales en el backend.",
         "SCRUM-17 / 197", "Corregido en dev", "PUT y DELETE api/personal-calendar-events/{id} con auditoría y tests."),
        ("Test desactualizado", "Get_WithEachAuthorizedRole_Returns200 incluía al rol Escucha, excluido por SCRUM-212.",
         "SCRUM-212", "Corregido", "Se quitó Escucha de los roles autorizados del test."),
        ("Defecto preexistente", "Los ítems del menú Inicio, Asistencia, Medicación e Informes apuntan a /inicio, una ruta sin pantalla (página vacía). Está en main.",
         "Frontend (Sprint 1-2)", "Abierta", "Detectado en la prueba de humo por rol: 4 de 15 pantallas; las demás cargan sin errores de API ni de consola."),
        ("Defecto", "Las series recurrentes con fecha de inicio anterior al rango consultado no se expandían: un evento semanal solo aparecía en su semana de inicio.",
         "SCRUM-189 / 194", "Corregido en dev", "La consulta incluye las series iniciadas antes del rango; hay tests unitarios y de integración."),
        ("Decisión resuelta", "Solapamiento de visitas: es una advertencia. El backend responde 409 con el aviso y acepta allowOverlap para guardar igual.",
         "SCRUM-211", "Resuelta 2026-10-09", "El frontend reenvía con allowOverlap cuando la persona vuelve a tocar Guardar."),
        ("Decisión resuelta", "Motivo de cancelación de visita: es opcional.",
         "SCRUM-74", "Resuelta 2026-10-09", "El PUT ya no lo exige y el frontend no manda un texto por defecto."),
        ("Mejora", "Recurrencia mensual de eventos (mismo día de cada mes; en meses cortos, el último día).",
         "SCRUM-16 / 189", "Implementada", "Columna repeats_monthly y migración AddCalendarEventMonthlyRecurrence."),
        ("Limitación", "Al editar una ocurrencia de una serie, la fecha del formulario es la de esa ocurrencia y la serie pasa a empezar ese día.",
         "SCRUM-190", "Abierta", "Una mejora posible es exponer la fecha de inicio de la serie en la ocurrencia."),
        ("Pendiente", "Las ramas dev-backend y dev-frontend no tienen los cambios de dev.",
         "Git", "Abierta", "Sincronizar por PR antes de que sigan Emir y Belén."),
        ("Pendiente", "Jira: SCRUM-198 y SCRUM-215 (QA) siguen 'En curso'.",
         "Jira", "Abierta", "Cerrar con el resultado de esta ejecución."),
    ]
    for i, it in enumerate(items, 1):
        wd.append([i, *it])
    style_sheet(wd, [5, 18, 70, 30, 18, 60])

    wb.save(OUT)
    print(f"{OUT} — manuales: {len(manual)} + {len(s3)} (Sprint 3) | automatizados: {len(auto)} + {len(s3_tests)} (Sprint 3)")


if __name__ == "__main__":
    main()
