"""Agrega al Word "Plan de prueba - completo formateado" los casos de la corrida del 16-17/10:
transversales de seguridad y carga (TC-195 a TC-211), pruebas de API del Sprint 4 (S4-A01..S4-M10) y los defectos #22 a #29.

Uso: python docs/qa/update_plan_docx.py <docx de entrada> <docx de salida>
Clona las tablas del último caso (TC-194) para conservar el formato; inserta todo antes de "Parte B".
"""
import copy
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

import run_2026_10_16 as run

HERE = Path(__file__).parent
PRE = "Stack local levantado (docker compose); usuarios de prueba *@test.com con los roles Referente, Directora, Coordinador y Escucha."


def set_cell(cell, text):
    """Reemplaza el texto de una celda conservando el formato del primer run."""
    paras = cell.paragraphs
    first = paras[0]
    for p in paras[1:]:
        p._element.getparent().remove(p._element)
    if first.runs:
        first.runs[0].text = text
        for r in first.runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        first.add_run(text)


def split_steps(text):
    rows = []
    for line in (text or "").splitlines():
        m = re.match(r"^(\d+)\.\s*(.*?)\s+→\s+(.*)$", line.strip())
        if m:
            rows.append((m.group(1), m.group(2), "—", m.group(3)))
        elif line.strip():
            rows.append((str(len(rows) + 1), line.strip(), "—", ""))
    return rows


def main():
    src, dst = sys.argv[1], sys.argv[2]
    doc = Document(src)
    body = doc.element.body
    kids = list(body.iterchildren())

    part_b = next(k for k in kids if k.tag.endswith("}p") and Paragraph(k, doc).text.startswith("Parte B"))
    h2_tpl = next(k for k in kids if k.tag.endswith("}p") and Paragraph(k, doc).text.startswith("Sprint 4 — Revisión mensual"))
    i194 = next(i for i, k in enumerate(kids) if k.tag.endswith("}tbl") and Table(k, doc).rows[0].cells[0].text.startswith("TC-194"))
    title_tpl, det_tpl, steps_tpl = kids[i194], kids[i194 + 1], kids[i194 + 2]
    blank_tpl = kids[i194 + 3]
    if any(Paragraph(k, doc).text.startswith("S4-A01") for k in kids if k.tag.endswith("}p")) or \
            any(Table(k, doc).rows[0].cells[0].text.startswith("TC-195") for k in kids if k.tag.endswith("}tbl")):
        sys.exit("El documento ya tiene los casos del 16-17/10; usar el documento de entrada anterior.")

    def heading(text):
        h = copy.deepcopy(h2_tpl)
        part_b.addprevious(h)
        Paragraph(h, doc).runs[0].text = text
        for r in Paragraph(h, doc).runs[1:]:
            r._element.getparent().remove(r._element)

    def case(tc_id, title, desc, pre, jira, test, steps):
        t = copy.deepcopy(title_tpl); d = copy.deepcopy(det_tpl); s = copy.deepcopy(steps_tpl); b = copy.deepcopy(blank_tpl)
        for el in (t, d, s, b):
            part_b.addprevious(el)
        set_cell(Table(t, doc).rows[0].cells[0], f"{tc_id} — {title}")
        dt = Table(d, doc)
        for row, val in zip(dt.rows, (desc, pre, jira, test)):
            set_cell(row.cells[1], val)
        st = Table(s, doc)
        data_tpl = st.rows[1]._tr
        for r in list(st.rows)[1:]:
            st._tbl.remove(r._tr)
        for n, step, dato, exp in steps:
            tr = copy.deepcopy(data_tpl)
            st._tbl.append(tr)
            cells = st.rows[-1].cells
            for c, v in zip(cells, (n, step, dato, exp)):
                set_cell(c, v)

    heading("Transversal — Seguridad básica y rendimiento (corrida del 2026-10-16)")
    for n, (sid, titulo, desc, pasos, script, real) in enumerate(run.SEC_CASES):
        case(f"TC-{run.SEC_FIRST_ID + n}", titulo, desc, PRE, "Transversal",
             f"Script {script} (caso {sid})", split_steps(pasos))

    results = HERE / "e2e" / "sprint4-api-results.json"
    if results.exists():
        heading("Sprint 4 — Pruebas de API de lo ya desarrollado: asistencia diaria (SCRUM-54 / 216) y esquema de medicación (SCRUM-59 / 232), 2026-10-17")
        for r in json.loads(results.read_text(encoding="utf-8")):
            att = r["id"].startswith("S4-A")
            call = "POST /api/attendances con el dato del caso" if att else "POST /api/persons/{id}/medication-schemes con el dato del caso"
            case(r["id"], r["name"], "Prueba contra la API real (sin interfaz).", PRE,
                 "SCRUM-54 / SCRUM-216" if att else "SCRUM-59 / SCRUM-232", "docs/qa/e2e/sprint4-api-flows.mjs",
                 [("1", "Iniciar sesión con el rol indicado en el título (Referente por defecto)", "—", "Se obtiene el token"),
                  ("2", call, "—", r["name"])])

    heading("Defectos hallados en la corrida del 2026-10-16 (#22 a #29)")
    for k, d in enumerate(run.NEW_DEFECTS):
        tipo, desc, hist, estado, nota, sev = d[:6]
        case(f"Hallazgo #{22 + k}", desc[:110] + ("…" if len(desc) > 110 else ""), desc, f"Severidad: {sev}. Tipo: {tipo}.", hist,
             f"Estado: {estado}", [("1", nota, "—", estado)])

    doc.save(dst)
    print("OK", dst)


if __name__ == "__main__":
    main()
